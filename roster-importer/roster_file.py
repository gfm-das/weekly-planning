"""Reading the transfer roster: the CSV or XLSX list of every missionary in the mission, with their area.

The office exports it from the Church's missionary system after each transfer. It must be COMPLETE: a missionary
who is not in it is treated as released. This file only reads the file into simple rows (one dict per missionary);
transfer.py compares them with the database and saves them.

The headings the file needs are in REQUIRED_COLUMNS. Only rows with the status Active or In-field are read.
"""
import csv
import io
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

REQUIRED_COLUMNS = ["Missionary", "ID", "Type", "Assignment", "Status", "Zone", "District", "Area", "Unit", "Email",
                    "Arrival Date", "Release Date", "Position", "Position Abbr"]
# Other names the export uses for a heading.
HEADER_ALIASES = {"Ecclesiastical Unit": "Unit", "Ecclesiastical Units": "Unit"}
# The position abbreviations that make someone a leader, and the leadership role each one gives.
ROLE_ALIASES = {"DL": "DL", "DLT": "DL", "ZL": "ZL", "ZL2": "ZL", "ZLL": "ZL", "STL": "STL", "STLL": "STL", "AP": "AP"}
ACTIVE_STATUSES = {"in-field", "in field", "active"}
DATE_FORMATS = ("%d %b %Y", "%Y-%m-%d", "%d.%m.%Y", "%m/%d/%Y")


def clean(value):
    """Text without extra spaces; None (and 0) become ''."""
    return " ".join(str(value or "").strip().split())


def parse_date(value):
    """A date from a cell: a spreadsheet date, or text like 04 Oct 2026, 2026-10-04, 04.10.2026 or 10/04/2026."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = clean(value)
    if not value:
        return None
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(value, date_format).date()
        except ValueError:
            pass
    raise ValueError(f"Unrecognized date: {value}")


def parse_name(raw, missionary_type):
    """('Tom', 'Smith', 'Elder Smith') from 'Smith, Tom' or 'Tom Smith'. Sisters are 'Sister Smith'."""
    raw = clean(raw)
    if "," in raw:
        last, rest = [part.strip() for part in raw.split(",", 1)]
        first = rest.split()[0] if rest else None
    else:
        parts = raw.split()
        first = parts[0] if parts else None
        last = parts[-1] if parts else raw
    title = "Sister" if clean(missionary_type).lower() == "sister" else "Elder"
    return first, last, f"{title} {last}"


def parse_units(raw):
    """The wards and branches of a row: 'Frankfurt 1st (66230), Frankfurt 2nd (English) (80144)' gives
    [{'name': 'Frankfurt 1st', 'unit_number': '66230'}, {'name': 'Frankfurt 2nd (English)', 'unit_number': '80144'}].
    A unit named twice is kept once."""
    raw = clean(raw)
    if not raw:
        return []
    units, seen = [], set()
    for part in [clean(x) for x in raw.split(",") if clean(x)]:
        numbered = re.match(r"^(.*)\s+\((\d+)\)\s*$", part)
        name, number = (clean(numbered.group(1)), numbered.group(2)) if numbered else (part, None)
        key = number or name.casefold()
        if key not in seen:
            seen.add(key)
            units.append({"name": name, "unit_number": number})
    return units


def roster_roles(row):
    """The leadership roles of a row from its Position Abbr ('ZL, DL' or 'ZL; DL'), each once, in file order."""
    tokens = [clean(x).upper() for x in (row.get("Position Abbr") or "").replace(";", ",").split(",") if clean(x)]
    found = []
    for token in tokens:
        role = ROLE_ALIASES.get(token)
        if role and role not in found:
            found.append(role)
    return found


def read_table(path):
    """All rows of the file (the first sheet of an XLSX) as lists of cell values."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        return list(csv.reader(io.StringIO(path.read_bytes().decode("utf-8-sig"))))
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = [list(r) for r in workbook[workbook.sheetnames[0]].iter_rows(values_only=True)]
        workbook.close()
        return rows
    raise ValueError("Upload a .csv or .xlsx roster.")


def iter_roster_rows(path):
    """(line number, {heading: value}) for every row under the heading row. Raises ValueError when a heading is
    missing."""
    rows = read_table(path)
    if not rows:
        raise ValueError("Roster file is empty.")
    headers = [HEADER_ALIASES.get(clean(h), clean(h)) for h in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(missing))
    for line_no, values in enumerate(rows[1:], start=2):
        yield line_no, {header: (values[i] if i < len(values) else None) for i, header in enumerate(headers) if header}


def read_roster(path):
    """The Active and In-field rows of a roster file as missionary dicts (see roster_row). Raises ValueError listing
    every problem (a missing or repeated ID, a missing zone, district or area, a date that cannot be read)."""
    rows, seen_ids, errors = [], set(), []
    for line_no, raw in iter_roster_rows(path):
        if clean(raw.get("Status")).lower() not in ACTIVE_STATUSES:
            continue
        number = clean(raw.get("ID"))
        if not number:
            errors.append(f"Row {line_no}: missing ID")
            continue
        if number in seen_ids:
            errors.append(f"Row {line_no}: duplicate ID {number}")
            continue
        seen_ids.add(number)
        if not clean(raw.get("Zone")) or not clean(raw.get("District")) or not clean(raw.get("Area")):
            errors.append(f"Row {line_no}: {number} is missing Zone, District, or Area")
            continue
        try:
            rows.append(roster_row(number, raw))
        except ValueError as e:
            errors.append(f"Row {line_no}: {e}")
    if errors:
        raise ValueError("\n".join(errors))
    if not rows:
        raise ValueError("No active/in-field roster rows found.")
    return rows


def roster_row(number, raw):
    """One missionary of the roster as the transfer uses it. Raises ValueError for a date that cannot be read."""
    arrival = parse_date(raw.get("Arrival Date"))
    release = parse_date(raw.get("Release Date"))
    first, last, display = parse_name(raw.get("Missionary"), raw.get("Type"))
    assignment = clean(raw.get("Assignment"))
    return {
        "missionary_number": number,
        "first_name": first,
        "last_name": last,
        "display_name": display,
        "missionary_type": clean(raw.get("Type")) or None,
        "arrival_date": arrival,
        "release_date": release,
        "zone": clean(raw.get("Zone")),
        "district": clean(raw.get("District")),
        "area": clean(raw.get("Area")),
        "units": parse_units(raw.get("Unit")),
        "email": clean(raw.get("Email")) or None,
        "position": clean(raw.get("Position")) or None,
        "position_abbr": clean(raw.get("Position Abbr")) or None,
        # Teaching is the usual assignment; anything else (Office, Visitor Center ...) is kept as a special one.
        "special_assignment": None if assignment.lower() in {"", "teaching"} else assignment,
        "roles": roster_roles(raw),
    }

"""Reading uploaded CSV and XLSX files, and the small helpers every upload type needs.

Nothing here touches the database or the disk: a file is read from the bytes of the upload, so the file itself
is never saved (the owner's rule for files with personal data).
"""
import csv
import io
import re
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from data_text import plain


class FileProblem(ValueError):
    """The file cannot be read at all (wrong type, empty, no heading row). Shown to the manager as it is."""


def clean(value):
    """Text without extra spaces. None and empty cells become ''."""
    return " ".join(str(value if value is not None else "").split())


def heading_key(text):
    """A heading in a form that ignores capitals, spaces and punctuation, so 'Size (km2)' matches 'size km2'."""
    return re.sub(r"[^a-z0-9]+", " ", clean(text).casefold()).strip()


def read_table(filename, data):
    """All rows of the first sheet (XLSX) or of the CSV, as lists of cell values. Empty rows are left out."""
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        rows = read_xlsx_rows(data)
    elif name.endswith(".csv") or name.endswith(".txt"):
        rows = read_csv_rows(data)
    else:
        raise FileProblem(plain("uploads.problem.wrongType"))
    rows = [row for row in rows if any(clean(v) for v in row)]
    if not rows:
        raise FileProblem(plain("uploads.problem.empty"))
    return rows


def read_xlsx_rows(data):
    # openpyxl is already part of DA Management (the roster import uses it). read_only keeps memory low.
    """The rows of the first sheet of an XLSX file, as lists of cell values."""
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:
        raise FileProblem(plain("uploads.problem.xlsx")) from error
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_csv_rows(data):
    # Exports from the Church and from Google Sheets are UTF-8; Excel on Windows may save in cp1252.
    """The rows of a CSV file, as lists of texts."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise FileProblem(plain("uploads.problem.encoding"))
    sample = text[:5000]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def find_heading_row(rows, required, look_at=10):
    """The first row (of the first `look_at`) that has every required heading. Some sheets have a title row first.

    Returns (row number, {heading key: column index}); the first column wins when a heading appears twice (the
    Church's finding export repeats several columns)."""
    wanted = [heading_key(h) for h in required]
    for number, row in enumerate(rows[:look_at]):
        columns = {}
        for index, cell in enumerate(row):
            columns.setdefault(heading_key(cell), index)
        if all(key in columns for key in wanted):
            return number, columns
    raise FileProblem(plain("uploads.problem.missingColumns", columns=", ".join(required)))


def cell(row, columns, *headings):
    """The value under the first of these headings that the file has, or None."""
    for heading in headings:
        index = columns.get(heading_key(heading))
        if index is not None and index < len(row):
            return row[index]
    return None


def read_date(value):
    """A date from a cell. Understands spreadsheet dates, 2026-09-27, 9/27/2026 (month first, as the Church's
    exports write it, with or without a time) and 27.09.2026 (day first, German). Empty gives None; anything
    else raises ValueError."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value)
    if not text:
        return None
    text = text.split(" ")[0]  # "9/27/2025 10:06:42 AM": only the day counts
    if "T" in text and re.match(r"^\d{4}-\d{2}-\d{2}T", text):
        text = text[:10]
    for pattern in ("%Y-%m-%d", "%m/%d/%Y", "%d.%m.%Y", "%m/%d/%y", "%d.%m.%y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError(plain("uploads.problem.notDate", value=clean(value)))


MONTH_WORDS = {  # English and German month names and the usual short forms (compared in small letters, without dots)
    "jan": 1, "january": 1, "januar": 1, "feb": 2, "february": 2, "februar": 2, "mar": 3, "march": 3, "mrz": 3, "maer": 3,
    "maerz": 3, "märz": 3, "mär": 3, "apr": 4, "april": 4, "may": 5, "mai": 5, "jun": 6, "june": 6, "juni": 6, "jul": 7,
    "july": 7, "juli": 7, "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9, "oct": 10, "okt": 10, "october": 10,
    "oktober": 10, "nov": 11, "november": 11, "dec": 12, "dez": 12, "december": 12, "dezember": 12}
NO_DATE_WORDS = {"-", "–", "—", "?", "??", "n/a", "na", "tba", "tbd", "none", "keine", "unbekannt", "unknown", "x", "xx"}
FIRST_DAY_NUMBER = 20000   # a spreadsheet date number: 20000 is 1954 and 60000 is 2064 (a number outside is not a date)
LAST_DAY_NUMBER = 60000


def loose_date(year, month, day):
    """A date from parts, with a two-digit year read as 20xx; ValueError (a plain 'not a date') when it does not exist."""
    year = year + 2000 if year < 100 else year
    return date(year, month, day)


def read_loose_date(value):
    """A date from a hand-typed cell (the office's Vollzogen sheets). Everything read_date reads, plus a spreadsheet date
    number, the English "Dec. 15th. 2025", "Jan. 4th, 2026", "15 December 2025", the German "15. Dezember 2025", "15 Dez 25",
    "15-12-2025" (day first) and "2025/12/15". A slash date is month first (9/27/2025, as the Church writes it) unless
    the first number cannot be a month (27/9/2025). Empty or a placeholder ("-", "?", "tba") gives None; anything else
    raises ValueError."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (datetime, date)):
        return read_date(value)
    if isinstance(value, (int, float)):
        if FIRST_DAY_NUMBER <= value <= LAST_DAY_NUMBER:
            return (datetime(1899, 12, 30) + timedelta(days=int(value))).date()
        raise ValueError(plain("uploads.problem.notDate", value=clean(value)))
    text = clean(value)
    if not text or text.casefold() in NO_DATE_WORDS:
        return None
    try:
        return read_date(text)
    except ValueError:
        pass
    bad = ValueError(plain("uploads.problem.notDate", value=text))
    try:
        found = re.fullmatch(r"(\d{4})[-/. ]+(\d{1,2})[-/. ]+(\d{1,2})", text)
        if found:
            return loose_date(int(found[1]), int(found[2]), int(found[3]))
        found = re.fullmatch(r"(\d{1,2})\s*([/.-])\s*(\d{1,2})\s*\2\s*(\d{2}|\d{4})", text)
        if found:
            first, sign, second, year = int(found[1]), found[2], int(found[3]), int(found[4])
            if sign == "/" and first <= 12:
                return loose_date(year, first, second)  # month first
            return loose_date(year, second, first)  # day first (German 15.12.2025, 15-12-2025, 27/9/2025)
        words = re.sub(r"(?<=\d)(st|nd|rd|th)\b", "", text.casefold())
        words = re.sub(r"[.,/]", " ", words).split()
        months = [MONTH_WORDS[w] for w in words if w in MONTH_WORDS]
        numbers = [w for w in words if w.isdigit()]
        if len(months) == 1 and len(numbers) == 2 and len(words) == 3:
            if len(numbers[0]) == 4:  # "2025 Dec 15"
                year, day = numbers[0], numbers[1]
            elif len(numbers[1]) == 4:  # "Dec 15 2025", "15 Dec 2025"
                day, year = numbers[0], numbers[1]
            else:  # "15 Dec 25": day, then a two-digit year
                day, year = numbers
            return loose_date(int(year), months[0], int(day))
    except ValueError:
        raise bad from None
    raise bad


YES_WORDS = {"x", "xx", "✓", "✔", "ja", "j", "yes", "y", "true", "1", "confirmed", "konfirmiert", "done"}
NO_WORDS = {"no", "nein", "n", "false", "0", "-", "–", "—", "?", "n/a"}


def read_confirmed(value):
    """The Konfirmiert cell of the Vollzogen sheet: (confirmed, date or None). An X (or Ja, Yes, a tick) is confirmed with
    no date; a date is confirmed on that day; empty, No or a dash is not confirmed. Anything else raises ValueError."""
    if value is None or isinstance(value, bool):
        return bool(value), None
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value), None
    if isinstance(value, (datetime, date, int, float)):
        day = read_loose_date(value)
        return (True, day) if day else (False, None)
    text = clean(value)
    if not text or text.casefold() in NO_WORDS:
        return False, None
    if text.casefold() in YES_WORDS:
        return True, None
    day = read_loose_date(text)
    return (True, day) if day else (False, None)


def read_count(value):
    """A whole number of zero or more, or None for an empty cell. Raises ValueError otherwise."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError(plain("uploads.problem.notNumber", value=value))
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = clean(value).replace(" ", "")
        if not text or text in {"-", "–"}:
            return None
        text = text.replace(",", ".") if re.fullmatch(r"-?\d+,\d+", text) else text.replace(",", "")
        try:
            number = float(text)
        except ValueError:
            raise ValueError(plain("uploads.problem.notNumber", value=clean(value))) from None
    if number < 0 or number != int(number):
        raise ValueError(plain("uploads.problem.notWhole", value=clean(value)))
    return int(number)


FC_NUMBER = re.compile(r"[A-Z0-9-]{3,30}")


def read_fc_number(value):
    """The FindeChristus (FC) number of a cell, written the one way every upload stores it: capitals, no spaces, and
    without a leading "FC" ("FC-123 456" and "123456" are both 123456). Leading zeros stay. None for an empty cell;
    ValueError for something that cannot be an FC number (it is 3 to 30 letters, digits or dashes)."""
    if value is None:
        return None
    if isinstance(value, float) and value == int(value):
        value = int(value)  # a number cell of a spreadsheet: 1234567.0 is 1234567
    text = clean(value).upper().replace(" ", "")
    text = re.sub(r"^FC[-#:.]*", "", text)
    if not text:
        return None if not clean(value) else _bad_fc(value)
    if not FC_NUMBER.fullmatch(text):
        _bad_fc(value)
    return text


def _bad_fc(value):
    raise ValueError(plain("uploads.problem.notFcNumber", value=clean(value)))


def read_big_count(value):
    """Like read_count, for numbers that can be large (a population): a text with thousands separators is read as the
    whole number it shows, so "21.000" (German) and "21,000" (English) are 21000, not 21. A spreadsheet number is as it is."""
    if isinstance(value, str):
        text = clean(value).replace(" ", "")
        if re.fullmatch(r"\d{1,3}(\.\d{3})+", text) or re.fullmatch(r"\d{1,3}(,\d{3})+", text):
            value = re.sub(r"[.,]", "", text)
    return read_count(value)


def read_decimal(value):
    """A number of zero or more that may have decimals (a size in km2), or None for an empty cell."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
    else:
        text = clean(value).replace(" ", "")
        if not text:
            return None
        # German 1.234,5 or 1.234.567 (a single dot, as in 0.135 or 1.234, is a decimal point)
        if re.fullmatch(r"\d{1,3}(\.\d{3})+,\d+", text) or re.fullmatch(r"\d{1,3}(\.\d{3}){2,}", text):
            text = text.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"\d+,\d+", text):  # 12,5
            text = text.replace(",", ".")
        else:
            text = text.replace(",", "")
        try:
            number = float(text)
        except ValueError:
            raise ValueError(plain("uploads.problem.notNumber", value=clean(value))) from None
    if number < 0:
        raise ValueError(plain("uploads.problem.belowZero", value=clean(value)))
    return number


def week_of(day):
    """The Sunday that ends the Monday-Sunday week this day is in (for event dates: a lesson on Monday 21 Sep
    counts for the week that ends on Sunday 27 Sep)."""
    return day + timedelta(days=6 - day.weekday())


def week_before(day):
    """The Sunday on or before this day (a Sunday stays itself). The zone history sheet names each week by its
    Sunday, so its dates are read with this."""
    return day - timedelta(days=(day.weekday() + 1) % 7)


def as_csv(headings, rows):
    """A CSV file (bytes, UTF-8 with the mark Excel needs to show umlauts) for a template download."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(headings)
    writer.writerows(rows)
    return ("﻿" + out.getvalue()).encode("utf-8")
